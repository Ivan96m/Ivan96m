from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

wb = Workbook()
ws = wb.active
ws.title = "Laplace 4x4"

# Matrix values
A = [
    [1, 2, 3, 4],
    [5, 6, 7, 8],
    [9, 1, 2, 3],
    [4, 5, 6, 7],
]

# Write title
ws["A1"] = "Macierz A (4x4)"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:D1")

# Write matrix starting row 2
start_row = 2
start_col = 1
for i, row in enumerate(A):
    for j, val in enumerate(row):
        cell = ws.cell(row=start_row + i, column=start_col + j, value=val)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EDEDED")

# Set column widths
for col in range(1, 9):
    ws.column_dimensions[get_column_letter(col)].width = 18

# Labels for minors
ws["F1"] = "Minory 3x3"
ws["F1"].font = Font(bold=True, size=12)
ws["F2"] = "M11"
ws["F3"] = "M12"
ws["F4"] = "M13"
ws["F5"] = "M14"
for cell in ["F2", "F3", "F4", "F5"]:
    ws[cell].font = Font(bold=True)

# Put formulas for minors in G2:G5
# Note: matrix is in A2:D5. We'll reference by those cells.
ws["G2"] = "=B3*(C4*D5-D4*C5)-C3*(B4*D5-D4*B5)+D3*(B4*C5-C4*B5)"  # M11
ws["G3"] = "=A3*(C4*D5-D4*C5)-C3*(A4*D5-D4*A5)+D3*(A4*C5-C4*A5)"  # M12
ws["G4"] = "=A3*(B4*D5-D4*B5)-B3*(A4*D5-D4*A5)+D3*(A4*B5-B4*A5)"  # M13
ws["G5"] = "=A3*(B4*C5-C4*B5)-B3*(A4*C5-C4*A5)+C3*(A4*B5-B4*A5)"  # M14

for cell in ["G2", "G3", "G4", "G5"]:
    ws[cell].number_format = "0"
    ws[cell].alignment = Alignment(horizontal="center")

# Laplace expansion
ws["F7"] = "Wyznacznik det(A) (Laplace po 1. wierszu)"
ws["F7"].font = Font(bold=True, size=12)
ws.merge_cells("F7:H7")

ws["F8"] = "Wzór"
ws["F8"].font = Font(bold=True)
ws["G8"] = "=A2*G2-B2*G3+C2*G4-D2*G5"
ws["G8"].alignment = Alignment(horizontal="center")

ws["F10"] = "Sprawdzenie (Excel)"
ws["F10"].font = Font(bold=True)
ws["G10"] = "=MDETERM(A2:D5)"
ws["G10"].alignment = Alignment(horizontal="center")

# Add borders around matrix and calculation area
thin = Side(style="thin", color="888888")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for r in range(2, 6):
    for c in range(1, 5):
        ws.cell(r, c).border = border

for r in range(2, 6):
    for c in range(6, 8):
        ws.cell(r, c).border = border

for r in [8, 10]:
    for c in range(6, 8):
        ws.cell(r, c).border = border

# Freeze panes
ws.freeze_panes = "A2"

file_path = "/mnt/data/wyznacznik_laplace_4x4_z_formulami.xlsx"
wb.save(file_path)

print(file_path)
